"""Experimento e concordância humana, offline (LLM falso)."""

import json

import pandas as pd
import pytest

from backend import run as run_module
from backend.experiments import baseline as baseline_module
from backend.experiments.experiment import run_experiment
from backend.experiments.human_agreement import agreement, cohen_kappa, validate, wilson
from tests.fakes import FakeLLM
from tests.test_graph import make_doc


def test_cohen_kappa_known_values():
    assert cohen_kappa(["a", "b", "a", "b"], ["a", "b", "a", "b"]) == 1.0
    assert cohen_kappa(["a", "a", "b", "b"], ["a", "b", "a", "b"]) == 0.0
    # exemplo clássico: po = 0,7, pe = 0,5 -> 0,4
    a = ["s"] * 5 + ["n"] * 5
    b = ["s"] * 4 + ["n"] + ["s"] * 2 + ["n"] * 3
    assert round(cohen_kappa(a, b), 3) == 0.4


def test_agreement_table():
    human = pd.DataFrame({"id": ["S01", "S02"], "nivel_percebido": ["Iniciante", "avançado"], "erro_factual": ["não", "sim"],
                          "tom_1a5": ["5", "3"], "trivializou": ["nao", "nao"]})
    machine = pd.DataFrame({"id": ["S01", "S02"], "nivel_previsto": ["iniciante", "intermediario"], "nivel_alvo": ["iniciante", "avancado"],
                            "erro_factual": ["nao", "sim"], "judge_tom": ["5", "5"], "trivializou": ["nao", "sim"]})
    t = agreement(human, machine).set_index("dimensao")
    assert t.loc["nivel (percebido x previsto)", "concordancia"] == 0.5
    assert t.loc["nivel (percebido x alvo)", "concordancia"] == 1.0
    assert t.loc["erro factual", "concordancia"] == 1.0
    assert t.loc["tom (erro absoluto médio, 1-5)", "concordancia"] == 1.0


def test_wilson_interval():
    lo, hi = wilson(9, 9)
    assert hi == 1.0 and 0.65 < lo < 0.71  # 9/9 ainda deixa ~30 p.p. de incerteza
    lo, hi = wilson(0, 9)
    assert lo == 0.0


def test_validate_flags_bad_rows():
    machine = pd.DataFrame({"id": ["S01", "S02", "S03"]})
    human = pd.DataFrame({
        "id": ["S01", "S02", "S04"],
        "nivel_percebido": ["básico", "", "Avançado"],
        "erro_factual": ["talvez", "", "sim"],
        "erro_descricao": ["", "", ""],
        "tom_1a5": ["7", "", "4"],
        "trivializou": ["N", "", "não"],
    })
    errors, warnings = validate(human, machine)
    text = " | ".join(errors)
    assert "S04" in text and "S03" in text  # id estranho e saída sem revisão
    assert "S01: nivel_percebido" in text and "S01: erro_factual" in text and "S01: tom_1a5" in text
    assert "S02: linha vazia" in text
    assert "S01: trivializou" not in text  # "N" vira "nao"
    assert "S04: erro_factual = sim sem erro_descricao" in warnings
    assert any(w.startswith("erro_factual preenchido em 2 de 3") for w in warnings)  # coluna opcional parcial


def test_optional_dimensions_left_blank_are_not_evaluated():
    machine = pd.DataFrame({"id": ["S01", "S02"], "nivel_previsto": ["iniciante", "avancado"], "nivel_alvo": ["iniciante", "avancado"],
                            "erro_factual": ["nao", "sim"], "judge_tom": ["5", "4"], "trivializou": ["nao", "nao"]})
    human = pd.DataFrame({"id": ["S01", "S02"], "nivel_percebido": ["Iniciante", "Avançado"], "erro_factual": ["", ""],
                          "erro_descricao": ["", ""], "tom_1a5": ["5", "4"], "trivializou": ["", ""]})
    errors, warnings = validate(human, machine)
    assert errors == [] and len(warnings) == 2
    t = agreement(human, machine).set_index("dimensao")
    assert t.loc["erro factual", "n"] == 0 and t.loc["nivel (percebido x alvo)", "concordancia"] == 1.0


def test_validate_accepts_clean_csv():
    machine = pd.DataFrame({"id": ["S01"]})
    human = pd.DataFrame({"id": ["S01"], "nivel_percebido": ["Intermediário"], "erro_factual": ["Não"],
                          "erro_descricao": [""], "tom_1a5": ["3"], "trivializou": ["sim"]})
    assert validate(human, machine) == ([], [])


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    for name in ("OUTPUTS_DIR", "RAW_DIR", "PROCESSED_DIR"):
        (tmp_path / name).mkdir()
        monkeypatch.setattr(run_module, name, tmp_path / name)
    monkeypatch.setattr(baseline_module, "OUTPUTS_DIR", tmp_path / "OUTPUTS_DIR")
    doc = make_doc()
    (tmp_path / "PROCESSED_DIR" / f"{doc.id}.json").write_text(doc.model_dump_json(), encoding="utf-8")
    return tmp_path


def test_experiment_end_to_end(workspace):
    llm = FakeLLM("fake:modelo")
    exp = run_experiment(["ata-teste"], [], "ata-teste", None, {"fast": llm, "strong": llm, "judge": llm},
                         workspace / "experiment", skip={"tradeoffs"})
    man = json.loads((exp / "manifest.json").read_text(encoding="utf-8"))
    assert set(man["hashes"]) == {"prompts_geracao", "prompts_avaliador", "prompt_baseline", "thresholds", "glossario"}
    metrics = pd.read_csv(exp / "csv" / "metricas.csv")
    assert set(metrics.system) == {"pipeline", "baseline"} and len(metrics) == 18
    var = pd.read_csv(exp / "csv" / "variancia.csv")
    assert var[var.tipo == "deterministica"].desvio_max.max() == 0
    human = pd.read_csv(exp / "revisao_humana.csv")
    assert len(human) == 9 and "nivel_alvo" not in human.columns  # revisão cega
    for fig in ("confusao_niveis", "distribuicao_metricas", "baseline_vs_pipeline", "reflection_antes_depois", "variancia"):
        assert (exp / "figuras" / f"{fig}.png").exists()
    report = (exp / "relatorio_experimental.md").read_text(encoding="utf-8")
    assert "SIMULAÇÃO" in report and "<!-- concordancia:inicio -->" in report


def test_resume_reuses_finished_documents_without_llm_calls(workspace):
    llm = FakeLLM("fake:modelo")
    llms = {"fast": llm, "strong": llm, "judge": llm}
    skip = {"tradeoffs", "variance", "human", "demo"}
    exp = run_experiment(["ata-teste"], [], None, None, llms, workspace / "experiment", skip=skip)
    first = pd.read_csv(exp / "csv" / "metricas.csv")
    assert (exp / "csv" / "parcial_ata-teste.json").exists()

    again = FakeLLM("fake:modelo")
    exp2 = run_experiment(["ata-teste"], [], None, None, {"fast": again, "strong": again, "judge": again},
                          workspace / "experiment", skip=skip, run_date=exp.name, resume=True)
    assert exp2 == exp and again.calls == []  # nada gerado nem avaliado de novo
    pd.testing.assert_frame_equal(pd.read_csv(exp / "csv" / "metricas.csv"), first)

    (exp / "csv" / "parcial_ata-teste.json").unlink()  # sem checkpoint: reaproveita runs/ e baseline/, só avalia
    third = FakeLLM("fake:modelo")
    run_experiment(["ata-teste"], [], None, None, {"fast": third, "strong": third, "judge": third},
                   workspace / "experiment", skip=skip, run_date=exp.name, resume=True)
    assert "ExtractionResult" not in third.calls and set(third.calls) <= {"ClaimCheck", "JudgeVerdict"}


def test_deterministic_comparison_reuses_saved_runs_without_llm(workspace):
    llm = FakeLLM("fake:modelo")
    skip = {"tradeoffs", "variance", "human", "demo"}
    exp = run_experiment(["ata-teste"], [], None, None, {"fast": llm, "strong": llm, "judge": llm}, workspace / "experiment", skip=skip)
    quiet = FakeLLM("fake:modelo")
    run_experiment(["ata-teste"], [], None, None, {"fast": quiet, "strong": quiet, "judge": quiet}, workspace / "experiment",
                   skip=skip, run_date=exp.name, llm_eval=False)
    assert quiet.calls == []
    metrics = pd.read_csv(exp / "csv" / "metricas.csv")
    assert set(metrics.system) == {"pipeline", "baseline"} and "judge_tom" not in metrics.columns
    assert json.loads((exp / "manifest.json").read_text(encoding="utf-8"))["comparacao"] == "deterministica"
    assert "só nas métricas determinísticas" in (exp / "relatorio_experimental.md").read_text(encoding="utf-8")


def test_tradeoffs_and_variance_reuse_saved_runs(workspace):
    llm = FakeLLM("fake:modelo")
    llms = {"fast": llm, "strong": llm, "judge": llm}
    base_skip = {"variance", "human", "demo"}
    exp = run_experiment(["ata-teste"], [], None, None, llms, workspace / "experiment", skip=base_skip | {"tradeoffs"})

    gen = FakeLLM("fake:modelo")
    run_experiment(["ata-teste"], [], None, None, {"fast": gen, "strong": gen, "judge": gen}, workspace / "experiment",
                   skip={"main", "human", "demo"}, run_date=exp.name, variance_repeats=2)
    trade = pd.read_csv(exp / "csv" / "tradeoffs.csv").set_index("config")
    assert list(trade.index) == ["barato + reflection", "barato sem reflection", "forte + reflection", "forte sem reflection"]
    assert (trade.loc[["barato sem reflection", "forte sem reflection"], "max_tentativas"] == 1).all()
    assert (trade.loc["barato sem reflection", "tokens"] <= trade.loc["barato + reflection", "tokens"])
    assert gen.calls.count("ExtractionResult") == 1  # só a execução forte é nova; a barata é a principal
    var = pd.read_csv(exp / "csv" / "variancia.csv")
    assert (var.repeticoes == 2).all() and var[var.tipo == "deterministica"].desvio_max.max() == 0
    assert (exp / "tradeoffs" / "ata-teste_forte.json").exists() and (exp / "csv" / "variancia_rep2.csv").exists()

    quiet = FakeLLM("fake:modelo")  # retomada: tudo vem dos arquivos salvos
    run_experiment(["ata-teste"], [], None, None, {"fast": quiet, "strong": quiet, "judge": quiet}, workspace / "experiment",
                   skip={"main", "human", "demo"}, run_date=exp.name, variance_repeats=2)
    assert quiet.calls == []


def test_human_sheet_roundtrip(tmp_path):
    from openpyxl import load_workbook

    from backend.experiments.human_sheet import SHEET, export, load_human

    pd.DataFrame({"id": ["S01"], "formato": ["texto"], "texto": ["Um texto."], "nivel_percebido": [""], "erro_factual": [""],
                  "erro_descricao": [""], "tom_1a5": [""], "trivializou": [""], "comentario": [""]}).to_csv(
        tmp_path / "revisao_humana.csv", index=False)
    path = export(tmp_path)
    wb = load_workbook(path)
    ws = wb[SHEET]
    for col, value in zip("DEGH", ("Intermediário", "não", 4, "sim"), strict=True):  # tom digitado como número
        ws[f"{col}2"] = value
    wb.save(path)
    human = load_human(tmp_path)
    assert human.loc[0, "tom_1a5"] == "4" and human.loc[0, "nivel_percebido"] == "Intermediário"
    assert validate(human, pd.DataFrame({"id": ["S01"]})) == ([], [])
