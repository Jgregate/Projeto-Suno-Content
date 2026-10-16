"""Planilha da revisão humana (revisao_humana.xlsx), mais fácil de preencher que o CSV.

    uv run python -m backend.experiments.human_sheet outputs/experiment/<data>

Gera a planilha a partir de revisao_humana.csv, com instruções, listas de opção e as colunas a preencher
destacadas. O backend.experiments.human_agreement lê a planilha quando ela existe.
"""

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SHEET = "Revisão"
# coluna do CSV -> (cabeçalho na planilha, largura, opções da lista ou None)
COLUMNS = {
    "id": ("ID", 7, None),
    "formato": ("Formato", 11, None),
    "texto": ("Texto gerado", 80, None),
    "nivel_percebido": ("Para quem o texto parece escrito?", 20, ["Iniciante", "Intermediário", "Avançado"]),
    "erro_factual": ("Tem erro factual?", 13, ["sim", "não"]),
    "erro_descricao": ("Se sim, qual é o erro?", 30, None),
    "tom_1a5": ("O tom combina com esse público? (1 a 5)", 18, ["1", "2", "3", "4", "5"]),
    "trivializou": ("Simplificou demais e perdeu algo importante?", 20, ["sim", "não"]),
    "comentario": ("Comentário (opcional)", 30, None),
}
TO_FILL = ["nivel_percebido", "erro_factual", "erro_descricao", "tom_1a5", "trivializou", "comentario"]

INSTRUCTIONS = [
    ("Revisão humana às cegas", True),
    ("", False),
    ("Na aba \"Revisão\" há 9 textos gerados pelo sistema a partir do release do Itaú (2T26).", False),
    ("Os textos estão embaralhados e sem o público pedido nem as notas do avaliador: leia cada um sem tentar adivinhar a intenção.", False),
    ("", False),
    ("Para cada linha, preencha as colunas amarelas (as que têm lista: clique na célula e use a seta):", False),
    ("• Para quem o texto parece escrito? Iniciante (nunca investiu), Intermediário (conhece o básico) ou Avançado (profissional/estudante de finanças).", False),
    ("• Tem erro factual? \"sim\" se algum número, data ou afirmação parece errado ou inventado; nesse caso, descreva na coluna seguinte.", False),
    ("• O tom combina com esse público? De 1 (nada adequado) a 5 (totalmente adequado), para o público que você marcou.", False),
    ("• Simplificou demais? \"sim\" se, para simplificar, o texto deixou de fora ou distorceu algo importante.", False),
    ("• Comentário: opcional.", False),
    ("", False),
    ("Não altere as colunas ID, Formato e Texto. Salve o arquivo no mesmo lugar, com o mesmo nome (revisao_humana.xlsx).", False),
]

YELLOW = PatternFill("solid", start_color="FFF6D5")
HEADER = PatternFill("solid", start_color="17181B")
THIN = Border(*(Side(style="thin", color="D0D0D0"),) * 4)


def export(exp: Path) -> Path:
    df = pd.read_csv(exp / "revisao_humana.csv", dtype=str).fillna("")
    wb = Workbook()
    guide = wb.active
    guide.title = "Como preencher"
    guide.column_dimensions["A"].width = 130
    for i, (text, bold) in enumerate(INSTRUCTIONS, start=1):
        cell = guide.cell(row=i, column=1, value=text)
        cell.font = Font(bold=bold, size=14 if bold else 11)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet(SHEET)
    keys = list(COLUMNS)
    for col, key in enumerate(keys, start=1):
        title, width, _ = COLUMNS[key]
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 48
    for row, rec in enumerate(df.to_dict("records"), start=2):
        for col, key in enumerate(keys, start=1):
            cell = ws.cell(row=row, column=col, value=rec.get(key, ""))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = THIN
            if key in TO_FILL:
                cell.fill = YELLOW
        ws.row_dimensions[row].height = min(409, 15 * (len(str(rec["texto"])) // 95 + str(rec["texto"]).count("\n") + 2))
    last = len(df) + 1
    for col, key in enumerate(keys, start=1):
        options = COLUMNS[key][2]
        if options:
            letter = ws.cell(row=1, column=col).column_letter
            dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
                                error="Escolha uma opção da lista.", errorTitle="Valor inválido", showErrorMessage=True)
            dv.add(f"{letter}2:{letter}{last}")
            ws.add_data_validation(dv)
    ws.freeze_panes = "D2"
    out = exp / "revisao_humana.xlsx"
    wb.save(out)
    return out


def load_human(exp: Path) -> pd.DataFrame:
    """Respostas da revisão: da planilha, se existir; senão do CSV. Colunas do CSV, valores como texto."""
    xlsx = exp / "revisao_humana.xlsx"
    if not xlsx.exists():
        return pd.read_csv(exp / "revisao_humana.csv", dtype=str).fillna("")
    ws = load_workbook(xlsx, data_only=True)[SHEET]
    rows = list(ws.iter_rows(values_only=True))
    by_title = {title: key for key, (title, _, _) in COLUMNS.items()}
    header = [by_title.get(str(h).strip(), str(h)) for h in rows[0]]
    records = [dict(zip(header, ("" if v is None else str(v).strip() for v in r), strict=False)) for r in rows[1:] if any(r)]
    df = pd.DataFrame(records, columns=list(COLUMNS)).fillna("")
    df["tom_1a5"] = df.tom_1a5.str.replace(r"\.0$", "", regex=True)  # número digitado vira 3.0 no openpyxl
    return df


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else sorted((Path("outputs") / "experiment").iterdir())[-1]
    print(f"Planilha gerada: {export(target)}")
